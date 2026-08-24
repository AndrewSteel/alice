"""
alice-mail-reader — HTTP wrapper around Python's imaplib.
Called by n8n (alice-mail-sync) to fetch email metadata and bodies.

Endpoints:
  GET  /health         — liveness check
  POST /encrypt        — encrypt a plaintext IMAP password (used by alice-mail-api)
  POST /test           — test IMAP login (accepts password_enc)
  POST /fetch          — fetch emails since a given IMAP UID, returns metadata + body preview
  POST /body           — fetch full body of a single email by UID
  POST /attachment     — fetch raw bytes of a single attachment by UID + attachment index
  POST /attachment-text — extract plaintext from a single PDF/DOCX/XLSX/ODT/ODS attachment by UID + index

Passwords are encrypted with AES-256-CBC. The key is derived from MAIL_ENC_KEY (env).
Plaintext passwords never leave this container.
"""

from __future__ import annotations

import base64
import email as email_lib
from email.header import decode_header
import hashlib
import imaplib
import io
import logging
import os
from datetime import datetime

import docx
import openpyxl
from odf.opendocument import load as odf_load
from odf.table import Table
from odf.table import TableCell as OdfTableCell
from odf.table import TableRow as OdfTableRow
from odf.text import P as OdfParagraph

from Crypto.Cipher import AES
from Crypto.Util.Padding import pad, unpad
from flask import Flask, request, jsonify
from pypdf import PdfReader

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("alice-mail-reader")

app = Flask(__name__)
MAX_EMAILS_PER_FETCH = 50

# Mirrors dms-extractor-pdf's PLAINTEXT_MAX_CHARS so both extraction paths feed
# downstream consumers the same maximum amount of text.
PLAINTEXT_MAX_CHARS = 50000


def _get_key() -> bytes:
    raw = os.environ.get("MAIL_ENC_KEY", "")
    if not raw:
        raise RuntimeError("MAIL_ENC_KEY not configured")
    return hashlib.sha256(raw.encode()).digest()


def _encrypt_password(plaintext: str) -> str:
    iv = os.urandom(16)
    cipher = AES.new(_get_key(), AES.MODE_CBC, iv)
    return iv.hex() + ":" + cipher.encrypt(pad(plaintext.encode("utf-8"), 16)).hex()


def _decrypt_password(password_enc: str) -> str:
    iv_hex, enc_hex = password_enc.split(":")
    cipher = AES.new(_get_key(), AES.MODE_CBC, bytes.fromhex(iv_hex))
    return unpad(cipher.decrypt(bytes.fromhex(enc_hex)), 16).decode("utf-8")


def _safe_decode(payload: bytes, charset: str | None) -> str:
    cs = charset or "utf-8"
    try:
        return payload.decode(cs, errors="replace")
    except (LookupError, UnicodeDecodeError):
        return payload.decode("latin-1", errors="replace")


def _decode_header(value: str | None) -> str:
    if not value:
        return ""
    parts = decode_header(value)
    out = []
    for raw, charset in parts:
        if isinstance(raw, bytes):
            out.append(_safe_decode(raw, charset))
        else:
            out.append(raw)
    return " ".join(out).strip()


def _get_text(msg) -> str:
    """Extract plaintext body (first text/plain part, max 500 chars)."""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if payload:
                    return _safe_decode(payload, part.get_content_charset())[:500]
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            return _safe_decode(payload, msg.get_content_charset())[:500]
    return ""


def _get_full_body(msg) -> str:
    """Extract full body, preferring HTML over plain-text."""
    html_body = ""
    plain_body = ""
    if msg.is_multipart():
        for part in msg.walk():
            ct = part.get_content_type()
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            text = _safe_decode(payload, part.get_content_charset())
            if ct == "text/html" and not html_body:
                html_body = text
            elif ct == "text/plain" and not plain_body:
                plain_body = text
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            plain_body = _safe_decode(payload, msg.get_content_charset())
    return html_body or plain_body


def _walk_attachment_parts(msg):
    """Yield (filename, part) for every attachment MIME part, in stable order.

    Single source of truth for attachment ordering: both the metadata list in
    /fetch and the index-based lookup in /attachment walk via this helper, so
    an attachment_index always refers to the same part.
    """
    for part in msg.walk():
        if part.get_content_maintype() == "multipart":
            continue
        if part.get("Content-Disposition") is None:
            continue
        filename = _decode_header(part.get_filename())
        if not filename:
            continue
        yield filename, part


def _get_attachments(msg) -> list[dict]:
    attachments = []
    for filename, part in _walk_attachment_parts(msg):
        payload = part.get_payload(decode=True)
        attachments.append({
            "name": filename,
            "mime_type": part.get_content_type(),
            "size_bytes": len(payload) if payload else 0,
        })
    return attachments


def _is_pdf(filename: str, mime_type: str) -> bool:
    return filename.lower().endswith(".pdf") or mime_type.lower() == "application/pdf"


# PROJ-91: DOCX/XLSX/ODT/ODS extension -> extractor mapping. Detection is by
# filename extension only (not MIME type) - unlike PDF, the MIME types email
# clients send for Office formats are inconsistent, while the extension is
# reliable and is already how dms-extractor-office routes files.
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".odt", ".ods"}


def _office_format(filename: str) -> str | None:
    name = filename.lower()
    for ext in OFFICE_EXTENSIONS:
        if name.endswith(ext):
            return ext
    return None


def _extract_docx_text(payload: bytes) -> str:
    document = docx.Document(io.BytesIO(payload))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_xlsx_text(payload: bytes) -> str:
    # read_only avoids loading the whole workbook into memory; data_only reads
    # the last-cached formula result instead of the formula text itself.
    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                cells = [str(cell.value) for cell in row if cell.value is not None]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    finally:
        workbook.close()


def _extract_odt_text(payload: bytes) -> str:
    doc = odf_load(io.BytesIO(payload))
    return "\n".join(str(p) for p in doc.getElementsByType(OdfParagraph))


def _extract_ods_text(payload: bytes) -> str:
    doc = odf_load(io.BytesIO(payload))
    lines = []
    for table in doc.getElementsByType(Table):
        for row in table.getElementsByType(OdfTableRow):
            cell_texts = []
            for cell in row.getElementsByType(OdfTableCell):
                text = "".join(str(p) for p in cell.getElementsByType(OdfParagraph))
                if text:
                    cell_texts.append(text)
            if cell_texts:
                lines.append("\t".join(cell_texts))
    return "\n".join(lines)


def _extract_office_text(payload: bytes, ext: str) -> str:
    """Extract plaintext from an Office document. Raises on any parse failure
    (caller catches and maps to status: extraction_failed, same as PDF)."""
    if ext == ".docx":
        return _extract_docx_text(payload)
    if ext == ".xlsx":
        return _extract_xlsx_text(payload)
    if ext == ".odt":
        return _extract_odt_text(payload)
    if ext == ".ods":
        return _extract_ods_text(payload)
    raise ValueError(f"unsupported office extension: {ext}")


def _extract_pdf_text(payload: bytes) -> tuple[str, int, bool]:
    """Extract plaintext from PDF bytes. Returns (text, page_count, truncated).

    Same contract as dms-extractor-pdf (pdf-parse + PLAINTEXT_MAX_CHARS cap),
    but synchronous and in-process instead of via MQTT/Redis.
    """
    reader = PdfReader(io.BytesIO(payload))
    page_count = len(reader.pages)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # A single broken page must not lose the text of all the others.
            parts.append("")
    text = "\n".join(parts)
    truncated = len(text) > PLAINTEXT_MAX_CHARS
    if truncated:
        text = text[:PLAINTEXT_MAX_CHARS]
    return text, page_count, truncated


def _fetch_attachment_part(data: dict, uid: str, attachment_index: int):
    """Fetch one attachment by absolute index. Returns (filename, mime, payload).

    Raises LookupError when the mail or the index does not exist. Shared by
    /attachment and /attachment-text so both address the exact same MIME part
    that /fetch reported at that position.
    """
    folder = data.get("folder", "INBOX")
    password = _decrypt_password(data["password_enc"])
    imap = _connect(data)
    imap.login(data["username"], password)
    imap.select(folder, readonly=True)

    typ, msg_data = imap.uid("fetch", uid.encode(), "(RFC822)")
    if typ != "OK" or not msg_data or msg_data[0] is None:
        imap.logout()
        raise LookupError("E-Mail nicht gefunden")

    raw = msg_data[0][1]
    msg = email_lib.message_from_bytes(raw)
    imap.logout()

    for idx, (filename, part) in enumerate(_walk_attachment_parts(msg)):
        if idx == attachment_index:
            return filename, part.get_content_type(), (part.get_payload(decode=True) or b"")

    raise LookupError("Anhang nicht gefunden")


def _connect(data: dict):
    host = data["host"]
    port = int(data.get("port", 993))
    ssl = data.get("ssl", True)
    if ssl:
        return imaplib.IMAP4_SSL(host, port)
    return imaplib.IMAP4(host, port)


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


@app.route("/encrypt", methods=["POST"])
def encrypt_password():
    data = request.json or {}
    plaintext = data.get("password", "")
    if not plaintext:
        return jsonify({"error": "password required"}), 400
    try:
        return jsonify({"password_enc": _encrypt_password(plaintext)})
    except Exception as exc:
        log.exception("encrypt_password failed")
        return jsonify({"error": str(exc)}), 500


@app.route("/test", methods=["POST"])
def test_connection():
    data = request.json or {}
    try:
        password = _decrypt_password(data["password_enc"])
        imap = _connect(data)
        imap.login(data["username"], password)
        imap.select("INBOX", readonly=True)
        imap.logout()
        return jsonify({"ok": True, "message": "Verbindung erfolgreich"})
    except imaplib.IMAP4.error as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    except Exception as exc:
        log.exception("test_connection failed")
        return jsonify({"ok": False, "message": str(exc)}), 500


@app.route("/fetch", methods=["POST"])
def fetch_emails():
    data = request.json or {}
    since_uid = int(data.get("since_uid", 0))
    since_date = data.get("since_date")  # ISO date string e.g. "2026-01-15", only used when since_uid=0
    limit = min(int(data.get("limit", 10)), MAX_EMAILS_PER_FETCH)
    folder = data.get("folder", "INBOX")

    try:
        password = _decrypt_password(data["password_enc"])
        imap = _connect(data)
        imap.login(data["username"], password)
        imap.select(folder, readonly=True)

        if since_uid > 0:
            typ, search_data = imap.uid("search", None, f"UID {since_uid + 1}:*")
        elif since_date:
            try:
                dt = datetime.strptime(since_date[:10], "%Y-%m-%d")
                imap_date = dt.strftime("%d-%b-%Y")  # IMAP format: 15-Jan-2026
                typ, search_data = imap.uid("search", None, f"SINCE {imap_date}")
            except (ValueError, Exception):
                typ, search_data = imap.uid("search", None, "ALL")
        else:
            typ, search_data = imap.uid("search", None, "ALL")

        if typ != "OK":
            imap.logout()
            return jsonify({"emails": [], "max_uid": since_uid, "error": "search failed"})

        uid_bytes = search_data[0].split() if search_data[0] else []
        uid_bytes = uid_bytes[:limit]

        emails = []
        max_uid = since_uid

        for uid_b in uid_bytes:
            uid_str = uid_b.decode()
            uid_int = int(uid_str)

            typ, msg_data = imap.uid("fetch", uid_b, "(RFC822)")
            if typ != "OK" or not msg_data or msg_data[0] is None:
                continue

            raw = msg_data[0][1]
            msg = email_lib.message_from_bytes(raw)

            attachments = _get_attachments(msg)
            emails.append({
                "uid": uid_int,
                "message_id": str(msg.get("Message-ID") or "").strip("<>").strip(),
                "subject": _decode_header(msg.get("Subject")),
                "sender": _decode_header(msg.get("From")),
                "recipients": _decode_header(msg.get("To")),
                "date": str(msg.get("Date") or ""),
                "body_preview": _get_text(msg),
                "attachments": attachments,
                "has_attachments": len(attachments) > 0,
            })
            max_uid = max(max_uid, uid_int)

        imap.logout()
        return jsonify({"emails": emails, "max_uid": max_uid})

    except imaplib.IMAP4.error as exc:
        return jsonify({"emails": [], "max_uid": since_uid, "error": str(exc)}), 400
    except Exception as exc:
        log.exception("fetch_emails failed")
        return jsonify({"emails": [], "max_uid": since_uid, "error": str(exc)}), 500


@app.route("/body", methods=["POST"])
def fetch_body():
    data = request.json or {}
    uid = str(data.get("uid", ""))
    folder = data.get("folder", "INBOX")

    if not uid:
        return jsonify({"error": "uid required"}), 400

    try:
        password = _decrypt_password(data["password_enc"])
        imap = _connect(data)
        imap.login(data["username"], password)
        imap.select(folder, readonly=True)

        typ, msg_data = imap.uid("fetch", uid.encode(), "(RFC822)")
        if typ != "OK" or not msg_data or msg_data[0] is None:
            imap.logout()
            return jsonify({"error": "E-Mail nicht gefunden"}), 404

        raw = msg_data[0][1]
        msg = email_lib.message_from_bytes(raw)
        body = _get_full_body(msg)
        imap.logout()
        return jsonify({"body": body})

    except imaplib.IMAP4.error as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        log.exception("fetch_body failed")
        return jsonify({"error": str(exc)}), 500


@app.route("/attachment", methods=["POST"])
def fetch_attachment():
    """Return the raw bytes of a single attachment (PROJ-53).

    Same connection fields as /body, plus attachment_index — a 0-based index
    into the attachment list produced by /fetch for the same UID.
    """
    data = request.json or {}
    uid = str(data.get("uid", ""))

    if not uid:
        return jsonify({"error": "uid required"}), 400

    try:
        attachment_index = int(data.get("attachment_index", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "attachment_index muss eine Zahl sein"}), 400

    if attachment_index < 0:
        return jsonify({"error": "Anhang nicht gefunden"}), 404

    try:
        filename, mime_type, payload = _fetch_attachment_part(data, uid, attachment_index)
        return jsonify({
            "filename": filename,
            "mime_type": mime_type,
            "size_bytes": len(payload),
            "content_base64": base64.b64encode(payload).decode("ascii"),
        })

    except LookupError as exc:
        return jsonify({"error": str(exc)}), 404
    except imaplib.IMAP4.error as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        log.exception("fetch_attachment failed")
        return jsonify({"error": str(exc)}), 500


@app.route("/attachment-text", methods=["POST"])
def fetch_attachment_text():
    """Return the extracted plaintext of a single PDF/DOCX/XLSX/ODT/ODS
    attachment (PROJ-53 iteration 2, extended by PROJ-91 for Office formats).

    Same input fields as /attachment. Used by the classification step of
    alice-mail-sync / alice-mail-attachment-backfill, which cannot run pdf-parse
    or Office parsers itself (n8n Code nodes only allow axios/redis/winston +
    crypto/fs/path).

    Unsupported attachment types and extraction failures return an empty text
    with a status field instead of an error, so the caller can fall back to
    filename + email context classification.
    """
    data = request.json or {}
    uid = str(data.get("uid", ""))

    if not uid:
        return jsonify({"error": "uid required"}), 400

    try:
        attachment_index = int(data.get("attachment_index", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "attachment_index muss eine Zahl sein"}), 400

    if attachment_index < 0:
        return jsonify({"error": "Anhang nicht gefunden"}), 404

    try:
        filename, mime_type, payload = _fetch_attachment_part(data, uid, attachment_index)

        if _is_pdf(filename, mime_type):
            try:
                text, page_count, truncated = _extract_pdf_text(payload)
            except Exception as exc:
                # Encrypted / malformed / image-only PDFs must not fail the caller.
                log.warning("PDF extraction failed for %s: %s", filename, exc)
                return jsonify({
                    "text": "",
                    "page_count": 0,
                    "truncated": False,
                    "status": "extraction_failed",
                })

            return jsonify({
                "text": text,
                "page_count": page_count,
                "truncated": truncated,
                "status": "ok",
            })

        office_ext = _office_format(filename)
        if office_ext:
            try:
                text = _extract_office_text(payload, office_ext)
            except Exception as exc:
                # Password-protected / corrupt / legacy-binary files must not fail the caller.
                log.warning("Office extraction failed for %s: %s", filename, exc)
                return jsonify({
                    "text": "",
                    "page_count": 0,
                    "truncated": False,
                    "status": "extraction_failed",
                })

            truncated = len(text) > PLAINTEXT_MAX_CHARS
            if truncated:
                text = text[:PLAINTEXT_MAX_CHARS]
            return jsonify({
                "text": text,
                "page_count": 0,
                "truncated": truncated,
                "status": "ok",
            })

        return jsonify({
            "text": "",
            "page_count": 0,
            "truncated": False,
            "status": "unsupported_format",
        })

    except LookupError as exc:
        return jsonify({"error": str(exc)}), 404
    except imaplib.IMAP4.error as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        log.exception("fetch_attachment_text failed")
        return jsonify({"error": str(exc)}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8007)
